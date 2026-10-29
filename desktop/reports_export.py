# reports_export.py — export/import/backup (fără Qt, ușor de testat)
import os
import csv
import shutil
import datetime
from typing import List, Dict, Any, Optional, Tuple

import database as db


def _backups_dir() -> str:
    """Folderul de backup pentru registrul activ, nu pentru vechiul cimitir.db."""
    from pathlib import Path
    import sys
    shared = str(Path(__file__).resolve().parent.parent / "shared")
    if shared not in sys.path:
        sys.path.insert(0, shared)
    import cimitir_backup
    import cimitir_config
    custom = (db.get_setting("backup_dir") or "").strip()
    d = Path(custom) if custom else cimitir_backup.default_backup_dir(cimitir_config.working_dir())
    d.mkdir(parents=True, exist_ok=True)
    return str(d)


def backup_db(dest_dir: Optional[str] = None) -> Optional[str]:
    """Backup consistent SQLite + media pentru registrul activ; compatibil cu bara MAIN."""
    try:
        from pathlib import Path
        import sys
        shared = str(Path(__file__).resolve().parent.parent / "shared")
        if shared not in sys.path:
            sys.path.insert(0, shared)
        import cimitir_backup
        import cimitir_config
        source = Path(cimitir_config.working_dir())
        target = cimitir_backup.backup(source, Path(dest_dir or _backups_dir()))
        problems = cimitir_backup.verify(target)
        if problems:
            raise RuntimeError("Verificarea backupului a eșuat: " + "; ".join(problems))
        return str(target)
    except Exception as e:
        print(f"[backup_db] {e}")
        return None


def restore_db(src_path: str) -> bool:
    """Nu înlocui o bază SQLite deschisă. Restaurarea sigură e în BackupWindow."""
    print("[restore_db] Restaurarea directă este dezactivată; folosește Copii de rezervă > Restaurează selecția.")
    return False


def auto_backup_daily() -> Optional[str]:
    """Backup zilnic prin mecanismul actual, inclusiv media registrului."""
    try:
        from pathlib import Path
        import sys
        shared = str(Path(__file__).resolve().parent.parent / "shared")
        if shared not in sys.path:
            sys.path.insert(0, shared)
        import cimitir_backup
        import cimitir_config
        target = cimitir_backup.run_scheduled(Path(cimitir_config.working_dir()), db.get_setting)
        return str(target) if target else None
    except Exception as e:
        print(f"[auto_backup_daily] {e}")
        return None


def export_dicts(rows: List[Dict[str, Any]], path: str,
                 headers: Optional[List[str]] = None) -> str:
    """Exportă o listă de dict-uri în XLSX (dacă e openpyxl) sau CSV (fallback)."""
    if not rows:
        rows = []
    if headers is None:
        headers = list(rows[0].keys()) if rows else []

    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(headers)
            for r in rows:
                ws.append([_cell(r.get(h)) for h in headers])
            wb.save(path)
            return path
        except ImportError:
            path = path[:-5] + ".csv"

    # CSV cu BOM (ca Excel să afișeze corect diacriticele)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow([_cell(r.get(h)) for h in headers])
    return path


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (list, dict)):
        return str(v)
    return v


def export_locuri(path: str) -> str:
    rows = db.list_locuri()
    headers = ["cod_loc", "sector", "parcela", "rand", "numar", "tip", "stare",
               "disponibilitate", "categorie", "suprafata", "capacitate", "nr_vechi", "numar_cf"]
    return export_dicts(rows, path, headers)


def export_inhumati(path: str) -> str:
    rows = db.load_all_inhumati()
    return export_dicts(rows, path)


def export_contracte(path: str) -> str:
    rows = db.load_contracte()
    return export_dicts(rows, path)


def _read_rows(path: str) -> List[Dict[str, Any]]:
    """Citește XLSX sau CSV → listă de dict-uri (prima linie = antet)."""
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            out.append({headers[i]: r[i] for i in range(len(headers))})
        return out
    # CSV
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def import_locuri(path: str) -> Tuple[int, int, List[str]]:
    """Import locuri din XLSX/CSV. Actualizează după cod_loc; creează rând gol dacă lipsește.
    Returnează (actualizate, create, erori)."""
    errors: List[str] = []
    updated = created = 0
    try:
        rows = _read_rows(path)
    except Exception as e:
        return 0, 0, [f"Nu pot citi fișierul: {e}"]

    allowed = ("sector", "parcela", "rand", "numar", "tip", "stare", "disponibilitate",
               "categorie", "suprafata", "capacitate", "nr_vechi", "numar_cf")
    conn = db.get_connection()
    try:
        for i, r in enumerate(rows, start=2):
            cod = (r.get("cod_loc") or r.get("Cod") or r.get("cod") or "")
            cod = str(cod).strip()
            if not cod:
                errors.append(f"Linia {i}: lipsește cod_loc")
                continue
            exists = conn.execute("SELECT 1 FROM locuri_de_veci WHERE cod_loc = ?", (cod,)).fetchone()
            if not exists:
                conn.execute("INSERT INTO locuri_de_veci (cod_loc, points) VALUES (?, '[]')", (cod,))
                created += 1
            else:
                updated += 1
            sets, vals = [], []
            for k in allowed:
                if k in r and r[k] not in (None, ""):
                    sets.append(f"{k} = ?")
                    vals.append(r[k])
            if sets:
                vals.append(cod)
                conn.execute(f"UPDATE locuri_de_veci SET {', '.join(sets)} WHERE cod_loc = ?", vals)
    finally:
        conn.close()
    return updated, created, errors
